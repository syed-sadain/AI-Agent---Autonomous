"""Word (.docx) rendering with python-docx: cover page, document control, contents,
styled sections/tables, header/footer page numbers, assumptions & generation notes."""
from __future__ import annotations

import re
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from app.agent.models import Block, Plan

NAVY = RGBColor(0x1F, 0x38, 0x64)
GREY = RGBColor(0x59, 0x59, 0x59)
CONTENT_WIDTH_CM = 16.6  # A4 (21cm) minus 2 x 2.2cm margins
_NUMERIC = re.compile(r"^[A-Z$€£₹]{0,3}\s?-?[\d,]+(\.\d+)?\s?(%|wk)?$")
_BOLD = re.compile(r"\*\*(.+?)\*\*")


def _shade(cell, hex_fill: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def _cell_margins(table, top=60, bottom=60, left=100, right=100) -> None:
    tblPr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for side, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tblPr.append(mar)


def _borders(table, color="BFC7D5") -> None:
    tblPr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), color)
        b.append(el)
    tblPr.append(b)


def _repeat_header(row) -> None:
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def _no_split(row) -> None:
    trPr = row._tr.get_or_add_trPr()
    trPr.append(OxmlElement("w:cantSplit"))


def _field(run, instr: str) -> None:
    for kind, text in (("begin", None), (None, instr), ("separate", None), (None, "1"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        elif text == instr:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = instr
        else:
            el = OxmlElement("w:t")
            el.text = text
        run._r.append(el)


def _rich(paragraph, text: str, *, size: float | None = None, color: RGBColor | None = None, bold=False) -> None:
    """Tiny inline-markdown: **bold** segments."""
    pos = 0
    for m in _BOLD.finditer(text):
        if m.start() > pos:
            _run(paragraph, text[pos : m.start()], size, color, bold)
        _run(paragraph, m.group(1), size, color, True)
        pos = m.end()
    if pos < len(text):
        _run(paragraph, text[pos:], size, color, bold)


def _run(paragraph, text, size=None, color=None, bold=False):
    r = paragraph.add_run(text)
    r.bold = bold or None
    if size:
        r.font.size = Pt(size)
    if color is not None:
        r.font.color.rgb = color
    return r


class DocumentBuilder:
    def build(
        self,
        plan: Plan,
        blocks: list[Block],
        out_path: Path,
        *,
        notes: list[str] | None = None,
        sources: list[str] | None = None,
    ) -> Path:
        doc = Document()
        self._setup(doc, plan)
        self._cover(doc, plan)
        self._document_control(doc, plan)
        self._contents(doc, blocks, has_appendix=True)

        for i, block in enumerate(blocks, 1):
            self._block(doc, i, block)

        self._appendix(doc, plan, notes or [], sources or [])
        self._header_footer(doc, plan)

        out_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(out_path)
        return out_path

    # ── setup ──────────────────────────────────────────────────────────────
    def _setup(self, doc, plan: Plan) -> None:
        sec = doc.sections[0]
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
        sec.left_margin = sec.right_margin = Cm(2.2)
        sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2.0)

        normal = doc.styles["Normal"]
        normal.font.name = "Calibri"
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        normal.font.size = Pt(11)
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing = 1.15

        for name, size, before in (("Heading 1", 17, 18), ("Heading 2", 13, 12)):
            st = doc.styles[name]
            st.font.name = "Calibri"
            st.font.size = Pt(size)
            st.font.bold = True
            st.font.color.rgb = NAVY
            st.paragraph_format.space_before = Pt(before)
            st.paragraph_format.space_after = Pt(6)
            st.paragraph_format.keep_with_next = True

        doc.core_properties.title = plan.title
        doc.core_properties.author = "Autonomous Document Agent"
        doc.core_properties.subject = plan.doc_type.replace("_", " ").title()

    # ── front matter ───────────────────────────────────────────────────────
    def _cover(self, doc, plan: Plan) -> None:
        for _ in range(5):
            doc.add_paragraph()
        label = doc.add_paragraph()
        _run(label, plan.doc_type.replace("_", " ").upper(), 11, GREY, True)
        title = doc.add_paragraph()
        title.paragraph_format.space_after = Pt(10)
        _run(title, plan.title, 30, NAVY, True)
        rule = doc.add_paragraph()
        pPr = rule._p.get_or_add_pPr()
        bdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        for k, v in (("val", "single"), ("sz", "12"), ("space", "1"), ("color", "1F3864")):
            bottom.set(qn(f"w:{k}"), v)
        bdr.append(bottom)
        pPr.append(bdr)
        obj = doc.add_paragraph()
        _run(obj, plan.objective, 13, GREY)
        doc.add_paragraph()
        meta = doc.add_paragraph()
        _run(meta, f"Prepared for: {plan.audience}", 11, GREY)
        meta.add_run().add_break()
        _run(meta, f"Date: {date.today().strftime('%d %B %Y')}", 11, GREY)
        meta.add_run().add_break()
        _run(meta, "Prepared by: Autonomous Document Agent (AI-generated draft)", 11, GREY)
        doc.add_page_break()

    def _document_control(self, doc, plan: Plan) -> None:
        doc.add_heading("Document Control", level=1)
        rows = [
            ["Document type", plan.doc_type.replace("_", " ").title()],
            ["Version", "1.0 (draft)"],
            ["Date", date.today().strftime("%d %B %Y")],
            ["Intended audience", plan.audience],
            ["Status", "AI-generated draft - requires human review before use"],
        ]
        self._table(doc, ["Field", "Value"], rows, 0)

    def _contents(self, doc, blocks: list[Block], has_appendix: bool) -> None:
        doc.add_heading("Contents", level=1)
        for i, b in enumerate(blocks, 1):
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            _run(p, f"{i}.  {b.heading}", 11, None, False)
        if has_appendix:
            p = doc.add_paragraph()
            _run(p, "Appendix.  Assumptions, Decisions & Generation Notes", 11, None, False)
        doc.add_page_break()

    # ── body ───────────────────────────────────────────────────────────────
    def _block(self, doc, number: int, block: Block) -> None:
        doc.add_heading(f"{number}. {block.heading}", level=1)
        if block.kind == "section":
            for para in block.paragraphs:
                _rich(doc.add_paragraph(), para)
            for bullet in block.bullets:
                _rich(doc.add_paragraph(style="List Bullet"), bullet.lstrip("-•* ").strip())
            if block.callout:
                self._callout(doc, block.callout)
        else:
            if block.intro:
                _rich(doc.add_paragraph(), block.intro)
            self._table(doc, block.columns, block.rows, block.bold_last_rows)
            if block.caption:
                cap = doc.add_paragraph()
                cap.paragraph_format.space_before = Pt(4)
                _run(cap, block.caption, 9, GREY, False).italic = True

    def _callout(self, doc, text: str) -> None:
        t = doc.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = t.rows[0].cells[0]
        _shade(cell, "EAF1FB")
        _cell_margins(t, 120, 120, 180, 180)
        p = cell.paragraphs[0]
        _run(p, "Key point  ", 10.5, NAVY, True)
        _rich(p, text, size=10.5)
        doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def _table(self, doc, columns: list[str], rows: list[list[str]], bold_last_rows: int):
        n = len(columns)
        table = doc.add_table(rows=1, cols=n)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        _borders(table)
        _cell_margins(table)

        # proportional column widths from content length
        weights = []
        for c in range(n):
            longest = max([len(columns[c])] + [len(r[c]) for r in rows])
            weights.append(min(max(longest, 8), 38))
        widths = [Cm(CONTENT_WIDTH_CM * w / sum(weights)) for w in weights]
        numeric = [
            all(_NUMERIC.match(r[c].strip()) for r in rows if r[c].strip()) and any(r[c].strip() for r in rows)
            for c in range(n)
        ]

        keep_together = len(rows) <= 12  # short tables stay on one page (totals never orphaned)
        hdr = table.rows[0]
        _repeat_header(hdr)
        for c, name in enumerate(columns):
            cell = hdr.cells[c]
            cell.width = widths[c]
            _shade(cell, "1F3864")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.keep_with_next = True
            _run(p, name, 10, RGBColor(0xFF, 0xFF, 0xFF), True)

        for ri, row in enumerate(rows):
            tr = table.add_row()
            _no_split(tr)
            is_total = bold_last_rows and ri >= len(rows) - bold_last_rows
            for c, text in enumerate(row):
                cell = tr.cells[c]
                cell.width = widths[c]
                if is_total:
                    _shade(cell, "DDE6F3")
                elif ri % 2 == 1:
                    _shade(cell, "F3F6FB")
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.keep_with_next = keep_together and ri < len(rows) - 1
                if numeric[c]:
                    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                _rich(p, text, size=10, bold=bool(is_total))
        doc.add_paragraph().paragraph_format.space_after = Pt(0)
        return table

    # ── appendix & chrome ──────────────────────────────────────────────────
    def _appendix(self, doc, plan: Plan, notes: list[str], sources: list[str]) -> None:
        doc.add_page_break()
        doc.add_heading("Appendix: Assumptions, Decisions & Generation Notes", level=1)
        doc.add_heading("Assumptions made by the agent", level=2)
        if plan.assumptions:
            for a in plan.assumptions:
                _rich(doc.add_paragraph(style="List Bullet"), a)
        else:
            doc.add_paragraph("The request was specific enough that no assumptions were required.")
        if plan.ambiguities:
            doc.add_heading("Ambiguities and conflicts identified", level=2)
            self._table(doc, ["Issue", "Decision taken"], [[a.issue, a.decision] for a in plan.ambiguities], 0)
        if sources:
            doc.add_heading("Reference data used", level=2)
            for s in sources:
                _rich(doc.add_paragraph(style="List Bullet"), s)
            _run(doc.add_paragraph(), "Benchmarks come from an illustrative mock knowledge base and must be validated.", 9, GREY, False).italic = True
        if notes:
            doc.add_heading("Generation warnings", level=2)
            for n in notes:
                _rich(doc.add_paragraph(style="List Bullet"), n)

    def _header_footer(self, doc, plan: Plan) -> None:
        sec = doc.sections[0]
        sec.different_first_page_header_footer = True  # clean cover page
        hp = sec.header.paragraphs[0]
        _run(hp, plan.title, 9, GREY, False)
        fp = sec.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        def styled():
            r = fp.add_run()
            r.font.size, r.font.color.rgb = Pt(9), GREY
            return r

        _run(fp, "AI-generated draft  |  Page ", 9, GREY, False)
        _field(styled(), "PAGE")
        _run(fp, " of ", 9, GREY, False)
        _field(styled(), "NUMPAGES")
