import io
import json
from pathlib import Path

from docx import Document
from fastapi.testclient import TestClient

from app.main import app

SAMPLES = Path(__file__).parent.parent / "samples"


def post(client, name):
    body = json.loads((SAMPLES / f"{name}_request.json").read_text())
    return client.post("/agent", json=body)


def test_health_lists_tools():
    with TestClient(app) as c:
        r = c.get("/health").json()
        assert r["status"] == "ok" and "assemble_document" in r["tools"]


def test_standard_request_end_to_end():
    with TestClient(app) as c:
        r = post(c, "standard")
        assert r.status_code == 200
        data = r.json()
        assert data["status"] == "completed" and data["plan"]["doc_type"] == "project_plan"
        assert len(data["tasks"]) >= 6 and all(t["status"] == "completed" for t in data["tasks"])
        docx = c.get(data["document"]["download_url"].replace("http://testserver", ""))
        assert docx.status_code == 200
        doc = Document(io.BytesIO(docx.content))
        text = "\n".join(p.text for p in doc.paragraphs)
        assert "Executive Summary" in text and len(doc.tables) >= 4
        assert "Total duration: 12 weeks" in text  # the "12-week" constraint reached the timeline tool


def test_complex_request_records_assumptions_and_conflicts():
    with TestClient(app) as c:
        data = post(c, "complex").json()
        assert len(data["plan"]["ambiguities"]) >= 3 and data["plan"]["assumptions"]
        issues = " ".join(a["issue"] for a in data["plan"]["ambiguities"]).lower()
        assert "budget" in issues and "build" in issues
        assert any(t["tool"] == "build_budget" for t in data["tasks"])


def test_guardrails_reject_bad_input():
    with TestClient(app) as c:
        assert c.post("/agent", json={"request": ""}).status_code == 422
        assert c.post("/agent", json={"request": "hi there ok"}).status_code == 422
        assert c.post("/agent", json={"request": "x" * 5000}).status_code == 422
        assert c.post("/agent", json={}).status_code == 422


def test_download_blocks_path_traversal_and_unknown_ids():
    with TestClient(app) as c:
        assert c.get("/documents/..%2F..%2Fetc%2Fpasswd").status_code in (400, 404)
        assert c.get("/documents/" + "0" * 32).status_code == 404


def test_base64_option():
    with TestClient(app) as c:
        body = json.loads((SAMPLES / "standard_request.json").read_text()) | {"include_base64": True}
        assert c.post("/agent", json=body).json()["document"]["base64"]
